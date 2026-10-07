# utils/pdf_parser.py
import re
import fitz  # PyMuPDF
import docx

def is_chart_label(text):
    """
    判断文本是否为文献图表中的零散标签（如 A, B, C, D 或者是孤立的图表简短坐标词）
    """
    cleaned = text.strip()
    
    # 1. 过滤单字母或双字母的子图编号（如 A, B, C, D, A B 等）
    if re.match(r"^[A-D](\s+[A-D])*$", cleaned):
        return True
        
    # 2. 过滤过短且不包含句号/标点的孤立字符（防止把图表里的坐标轴单位或残缺词当成正文）
    if len(cleaned) <= 3 and not any(p in cleaned for p in [".", ",", ";", ":", "—"]):
        # 保留常见的医学缩写，过滤纯粹的图表噪点
        if cleaned.lower() not in ["ph", "co2", "o2", "iv", "im"]:
            return True
            
    return False

def extract_structured_content(uploaded_file, start_page=1, end_page=None):
    """
    优化版解析器：过滤图表零散英文标签，确保只提取连贯的正文段落进行翻译
    """
    file_name = uploaded_file.name
    file_extension = file_name.split(".")[-1].lower()
    
    elements = []
    
    if file_extension == "pdf":
        pdf_bytes = uploaded_file.read()
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        total_pages = len(doc)
        
        if end_page is None or end_page > total_pages:
            end_page = total_pages
            
        start_idx = max(0, start_page - 1)
        end_idx = min(total_pages, end_page)
        
        for page_num in range(start_idx, end_idx):
            page = doc[page_num]
            elements.append({"type": "text", "content": f"\n\n## Page {page_num + 1}\n"})
            
            blocks = page.get_text("blocks")
            sorted_blocks = sorted(blocks, key=lambda b: (b[1], b[0]))
            
            for b in sorted_blocks:
                block_type = b[6] if len(b) > 6 else 0
                if block_type == 0:  # 纯文本块
                    text = b[4].strip()
                    if text:
                        # 如果是图表零散标签，直接跳过，不让其进入翻译队列
                        if is_chart_label(text):
                            continue
                        elements.append({"type": "text", "content": text})
                        
        return elements, total_pages
        
    elif file_extension in ["docx", "doc"]:
        doc = docx.Document(uploaded_file)
        for p in doc.paragraphs:
            text = p.text.strip()
            if text and not is_chart_label(text):
                elements.append({"type": "text", "content": text})
        return elements, len(doc.paragraphs)
        
    elif file_extension in ["md", "txt"]:
        text_content = uploaded_file.read().decode("utf-8", errors="ignore")
        for line in text_content.split("\n"):
            text = line.strip()
            if text and not is_chart_label(text):
                elements.append({"type": "text", "content": text})
        return elements, len(text_content.split("\n"))
        
    return [], 0
